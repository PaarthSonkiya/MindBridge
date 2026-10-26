#!/usr/bin/env python3
"""Fill the organisers' Prototype Submission template with MindBridge's content and measured results.

    python docs/submission/build_submission.py --template template.docx \
        --out "MindBridge_Prototype Submission_V3.docx" [--board bench/results/<board-voice-run>]

Text lives in docs/submission/content.py. Numbers are read from the published result files, so the
document always matches the evidence. Anything that can only be measured on the team's own RB3 Gen 2
(or recorded on video) stays as a yellow-highlighted "[TO FILL …]" until `--board` points at a voice
benchmark run made on the board.
"""

from __future__ import annotations

import argparse
import copy
import importlib.util
import json
import re
import sys
from pathlib import Path
from types import SimpleNamespace

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
INK = RGBColor(0x1A, 0x1A, 0x2E)
TODO = re.compile(r"(\[TO FILL[^\]]*\])")


def _load_content():
    spec = importlib.util.spec_from_file_location("content", HERE / "content.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _font_of(run):
    return SimpleNamespace(name=run.font.name, size=run.font.size)


def _style_answer(run, font):
    """Answer text: same font/size as the template's placeholder, upright, ink colour."""
    run.font.name = font.name
    run.font.size = font.size
    run.italic = False
    run.font.color.rgb = INK


def _add_text(paragraph, text: str, font, before=None):
    """Append `text` to a paragraph; **bold** spans and [TO FILL …] (highlighted) are recognised."""
    runs = []
    for part in re.split(r"(\*\*[^*]+\*\*|\[TO FILL[^\]]*\])", text):
        if not part:
            continue
        r = paragraph.add_run()
        if before is not None:  # move the new run to just before `before`
            before._r.addprevious(r._r)
        _style_answer(r, font)
        if part.startswith("**"):
            r.text = part[2:-2]
            r.bold = True
        else:
            r.text = part
            if TODO.fullmatch(part):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                r.bold = True
        runs.append(r)
    return runs


def _placeholder_runs(paragraph):
    return [r for r in paragraph.runs if r.italic and r.text.strip().startswith("[")]


def fill_inline(paragraph, answers: list[str]):
    """Replace each italic [placeholder] run in order; keeps a trailing line break if it had one."""
    ph = _placeholder_runs(paragraph)
    assert len(ph) == len(answers), (paragraph.text, len(ph), len(answers))
    for run, ans in zip(ph, answers, strict=True):
        tail = "\n" if run.text.endswith("\n") else ""
        new = _add_text(paragraph, ans + tail, _font_of(run), before=run)
        run._r.getparent().remove(run._r)
        assert new


def fill_cell(cell, text: str):
    """Replace a table cell's placeholder with one or more paragraphs (split on blank lines)."""
    p0 = cell.paragraphs[0]
    ph = _placeholder_runs(p0)[0]
    font = _font_of(ph)
    blocks = [b.strip() for b in text.strip().split("\n\n")]
    _add_text(p0, blocks[0], font, before=ph)
    ph._r.getparent().remove(ph._r)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for b in blocks[1:]:
        newp = copy.deepcopy(p0._p)
        for child in list(newp):
            if child.tag != qn("w:pPr"):
                newp.remove(child)
        cell._tc.append(newp)
        _add_text(cell.paragraphs[-1], b, font)


def find_para(doc, startswith: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise KeyError(startswith)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--board", default="", help="folder of a voice benchmark run made on the RB3 Gen 2")
    ap.add_argument("--manual", default="", help="JSON with values only a person can supply (video link, "
                                                 "package link, test conditions…); see manual_fields.example.json")
    args = ap.parse_args()

    C = _load_content()
    board = json.loads((Path(args.board) / "summary.json").read_text()) if args.board else None
    F = C.fields(board)
    if args.manual:
        for key, value in json.loads(Path(args.manual).read_text(encoding="utf-8")).items():
            if key.startswith("_") or not value:
                continue
            if isinstance(value, dict):
                for sub, text in value.items():
                    if key == "results" and any(text):  # [result text, evidence]; keep ours where empty
                        F[key][sub] = tuple(t or old for t, old in zip(text, F[key][sub], strict=True))
                    elif key != "results" and text:
                        F[key][sub] = text
            else:
                F[key] = value
    doc = docx.Document(args.template)

    # Section 1 table
    t1 = doc.tables[0]
    for row in t1.rows[1:]:
        label = row.cells[0].text.strip()
        fill_cell(row.cells[1], F["summary"][label])

    # Section 2
    fill_inline(find_para(doc, "Video link:"), [F["video_link"], F["video_duration"]])
    fill_inline(find_para(doc, "Key timestamps:"), [F["video_timestamps"]])

    # Section 3.1
    fill_inline(find_para(doc, "Offline package:"), [F["package_url"], F["package_format"], F["source_repo"]])
    fill_inline(find_para(doc, "Package contents:"), [F["package_contents"]])

    # Section 3.2: diagram + notes
    p = find_para(doc, "[Insert a compact diagram")
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.add_run().add_picture(str(ROOT / "docs" / "img" / "architecture_mac.png"), width=Cm(16.5))
    fill_inline(find_para(doc, "Implementation notes:"), [F["implementation_notes"]])

    # Section 3.3 table
    t2 = doc.tables[1]
    for row in t2.rows[1:]:
        fill_cell(row.cells[1], F["requirements"][row.cells[0].text.strip()])

    # Section 3.4 steps
    for label, text in F["steps"].items():
        fill_inline(find_para(doc, label), [text])

    # Section 4
    fill_inline(find_para(doc, "Test record:"), [F["test_record"], F["evidence_folder"]])
    t3 = doc.tables[2]
    for row in t3.rows[1:]:
        area = row.cells[0].text.strip()
        res, ev = F["results"][area]
        fill_cell(row.cells[1], res)
        fill_cell(row.cells[2], ev)
    fill_inline(find_para(doc, "Known limitations:"), [F["limitations"]])
    fill_inline(find_para(doc, "Next development steps:"), [F["next_steps"]])

    doc.core_properties.title = "MindBridge — Prototype Submission (Phase 3)"
    doc.core_properties.author = "MindBridge"
    doc.save(args.out)
    left = sum(len(TODO.findall(p.text)) for p in doc.paragraphs) + sum(
        len(TODO.findall(c.text)) for t in doc.tables for row in t.rows for c in row.cells)
    print(f"wrote {args.out}; {left} [TO FILL] items remain")


if __name__ == "__main__":
    sys.exit(main())
