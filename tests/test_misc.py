import socket

import numpy as np
import pytest


def test_offline_guard_blocks_external_allows_loopback():
    from mindbridge import offline_guard

    offline_guard.install()
    with pytest.raises(offline_guard.NetworkBlocked):
        socket.create_connection(("93.184.216.34", 80), timeout=1)
    srv = socket.socket()
    srv.bind(("127.0.0.1", 0))
    srv.listen(1)
    c = socket.create_connection(srv.getsockname(), timeout=1)
    c.close()
    srv.close()


def test_mel_matches_reference_when_available():
    from mindbridge.asr.mel import log_mel

    x = (np.random.RandomState(0).randn(16000 * 2) * 0.1).astype(np.float32)
    mine = log_mel(x)
    assert mine.shape == (1, 80, 3000) and mine.dtype == np.float32
    tf = pytest.importorskip("transformers")
    ref = tf.WhisperFeatureExtractor(feature_size=80)(x, sampling_rate=16000, return_tensors="np")["input_features"]
    assert np.abs(ref - mine).max() < 1e-4


def test_keep_final_punctuation():
    from mindbridge.mt.text import keep_final_punctuation

    assert keep_final_punctuation("Tôi muốn đi.", "가고 싶어요", "ko") == "가고 싶어요."
    assert keep_final_punctuation("Anh đi đâu?", "你去哪里，", "zh") == "你去哪里？"
    assert keep_final_punctuation("Xin chào", "Hello", "en") == "Hello"
    assert keep_final_punctuation("Xin chào.", "Hello!", "en") == "Hello!"


def test_chinese_punctuation():
    from mindbridge.mt.text import chinese_punctuation

    assert chinese_punctuation("让我先介绍一下,这是我的校长.") == "让我先介绍一下，这是我的校长。"
    assert chinese_punctuation("你看过我的护照吗?") == "你看过我的护照吗？"
    assert chinese_punctuation("价格是3.5元. OK, fine.") == "价格是3.5元。OK, fine."


def test_sentence_split_and_simplified_chinese():
    from mindbridge.mt.text import clean_asr, split_sentences

    assert split_sentences("Stop. Wait here! Ok?") == ["Stop.", "Wait here!", "Ok?"]
    assert split_sentences("停机。等一下！") == ["停机。", "等一下！"]
    out = clean_asr("在更換壓力閥之前", "zh")
    assert "换" in out or "更換" not in out  # converted when opencc is installed


def test_config_env_override(monkeypatch):
    from mindbridge.config import load_config

    monkeypatch.setenv("MINDBRIDGE_ASR_BACKEND", "cpu")
    monkeypatch.setenv("MINDBRIDGE_PORT", "9000")
    cfg = load_config()
    assert cfg.asr.backend == "cpu" and cfg.port == 9000


def test_bench_loader(tmp_path):
    import importlib.util
    from pathlib import Path

    spec = importlib.util.spec_from_file_location("rb", Path(__file__).parent.parent / "bench" / "run_benchmark.py")
    rb = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(rb)
    (tmp_path / "a.csv").write_text("English,Vietnamese\nStop the machine.,Dừng máy lại.\n", encoding="utf-8")
    (tmp_path / "vi-ko.tsv").write_text("src\ttgt\nDừng máy.\t기계를 멈추세요.\n", encoding="utf-8")
    pairs = rb.load_pairs(tmp_path)
    assert pairs[("vi", "en")][0] == {"a": "Dừng máy lại.", "b": "Stop the machine.", "file": "a.csv"}
    assert pairs[("vi", "ko")][0]["b"] == "기계를 멈추세요."
    assert rb.error_rate("dừng máy lại", "dừng máy", "vi") == pytest.approx(1 / 3)


def test_bench_loader_audio_columns_and_sheets(tmp_path):
    import importlib.util
    from pathlib import Path

    import openpyxl

    spec = importlib.util.spec_from_file_location("rb2", Path(__file__).parent.parent / "bench" / "run_benchmark.py")
    rb = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(rb)
    (tmp_path / "t.csv").write_text("id,audio_vi,vi,ko\n1,a/1.wav,Dừng máy.,기계를 멈추세요.\n", encoding="utf-8")
    wb = openpyxl.Workbook()
    wb.active.append(["Vietnamese", "English"])
    wb.active.append(["Dừng máy.", "Stop the machine."])
    ws2 = wb.create_sheet("ko")
    ws2.append(["Vietnamese", "Korean"])
    ws2.append(["Dừng máy.", "기계를 멈추세요."])
    wb.save(tmp_path / "s.xlsx")
    pairs = rb.load_pairs(tmp_path)
    ko = pairs[("vi", "ko")]
    assert any(i["a"] == "Dừng máy." and i.get("audio_a", "").endswith("1.wav") for i in ko)
    assert len(ko) == 2 and pairs[("vi", "en")][0]["b"] == "Stop the machine."


def _load_bench(name):
    import importlib.util
    from pathlib import Path

    spec = importlib.util.spec_from_file_location(name, Path(__file__).parent.parent / "bench" / "run_benchmark.py")
    rb = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(rb)
    return rb


def test_bench_loader_organiser_layout(tmp_path):
    """The organisers' sample set: xlsx + wav/<ID>.wav, and <Lang>_text/<ID>.txt + <Lang>_voice/<ID>.wav."""
    import openpyxl

    rb = _load_bench("rb3")
    ve = tmp_path / "Vie-Eng dataset" / "Vie-Eng"
    (ve / "wav").mkdir(parents=True)
    wb = openpyxl.Workbook()
    wb.active.append(["ID", "Tiếng Việt", "Tiếng Anh - Final"])
    wb.active.append(["VE-0001", "Dừng máy lại.", "Stop the machine."])
    wb.save(ve / "OneVoice_V_to_E_samples_100.xlsx")
    (ve / "wav" / "VE-0001.wav").write_bytes(b"RIFF")
    ev = tmp_path / "Vie-Eng dataset" / "Eng-Vie"
    (ev / "wav").mkdir(parents=True)
    wb.save(ev / "OneVoice_V_to_E_samples_100.xlsx")  # the same sheet copied by mistake
    (ev / "wav" / "EV-0002.wav").write_bytes(b"RIFF")
    vk = tmp_path / "Vie-Kor dataset"
    for d in ("Kor_text", "Vie_text", "Kor_voice", "Vie_voice"):
        (vk / d).mkdir(parents=True)
    (vk / "Kor_text" / "VK-0003.txt").write_text("기계를 멈추세요.\n", encoding="utf-8")
    (vk / "Vie_text" / "VK-0003.txt").write_text("Dừng máy.", encoding="utf-8")
    (vk / "Kor_voice" / "VK-0003.wav").write_bytes(b"RIFF")

    pairs, unlinked = rb.load_dataset(tmp_path)
    en = pairs[("vi", "en")]
    assert len(en) == 1 and en[0]["id"] == "VE-0001" and en[0]["native"] == ["vi-en"]
    assert en[0]["audio_a"].endswith("VE-0001.wav") and "audio_b" not in en[0]
    ko = pairs[("vi", "ko")]
    assert ko[0]["a"] == "Dừng máy." and ko[0]["b"] == "기계를 멈추세요."
    assert ko[0]["audio_b"].endswith("Kor_voice/VK-0003.wav") and ko[0]["native"] == ["ko-vi"]
    assert [(u["id"], u["src"], u["tgt"]) for u in unlinked] == [("EV-0002", "en", "vi")]
    assert rb._lang_of_column("Tiếng Anh - Final") == "en"
    assert rb._langs_in_name("OneVoice_V_to_E_samples_100") == ["vi", "en"]
    assert rb._langs_in_name("Vie-Man dataset") == ["vi", "zh"] and rb._id_langs("VM-0223") == ["vi", "zh"]


def test_envit5_english_normalisation():
    from mindbridge.mt.envit5_ct2 import normalise_english

    assert normalise_english("Do not restart Machine 4 until the valve is replaced.") == \
        "Do not restart machine 4 until the valve is replaced."
    assert normalise_english("Run Line 2 at half speed; check Pump #3.") == \
        "Run production line 2 at half speed; check pump #3."
    assert normalise_english("Machine learning is fun.") == "Machine learning is fun."
    from mindbridge.mt.envit5_ct2 import detokenise_english

    assert detokenise_english("Do n't add sugar. I 've seen it, is n't it?") == "Don't add sugar. I've seen it, isn't it?"
    from mindbridge.mt.envit5_ct2 import tidy_vietnamese

    assert tidy_vietnamese("máy 6 dừng lúc 21: 55.") == "Máy 6 dừng lúc 21:55."


def test_mt_router_picks_backend(tmp_path):
    from mindbridge.config import MtConfig
    from mindbridge.mt import MtRouter

    d = tmp_path / "mt" / "envit5-translation-ct2-int8"
    d.mkdir(parents=True)
    (d / "model.bin").write_bytes(b"")
    r = MtRouter(tmp_path, MtConfig())
    assert r.backend_key("vi", "en") == "envit5" and r.backend_key("en", "vi") == "envit5"
    assert r.backend_key("vi", "ko") == "nllb" and r.backend_key("en", "zh") == "nllb"
    r2 = MtRouter(tmp_path, MtConfig(en_vi_model_dir=""))
    assert r2.backend_key("vi", "en") == "nllb"


def test_macos_route_parsing():
    from mindbridge.metrics import parse_darwin_route

    online = ("   route to: default\ndestination: default\n       mask: default\n    gateway: 192.168.1.1\n"
              "  interface: en0\n      flags: <UP,GATEWAY,DONE,STATIC,PRCLONING,GLOBAL>\n")
    tunnel_v6 = "   route to: ::\ndestination: default\n    gateway: fe80::%utun0\n  interface: utun0\n"
    assert parse_darwin_route(online) == "default via en0 gateway 192.168.1.1 (IPv4)"
    assert parse_darwin_route("", tunnel_v6) == ""  # Wi-Fi off: only the system tunnel's link-local route
    assert parse_darwin_route("", "") == ""


def test_peak_memory_units(monkeypatch):
    import resource

    from mindbridge import metrics

    class R:
        ru_maxrss = 2 * 2**30  # 2 GiB in bytes on macOS, would be KiB on Linux

    monkeypatch.setattr(resource, "getrusage", lambda _who: R)
    monkeypatch.setattr(metrics, "IS_MAC", True)
    assert round(metrics.peak_rss_mb()) == 2048
    monkeypatch.setattr(metrics, "IS_MAC", False)
    R.ru_maxrss = 2 * 2**20  # 2 GiB in KiB
    assert round(metrics.peak_rss_mb()) == 2048


def test_locks_have_websocket_support():
    """uvicorn serves the UI's /ws only with a WebSocket library installed (the UI is blank without it)."""
    from pathlib import Path

    root = Path(__file__).parent.parent
    for f in ("requirements.txt", "requirements-aarch64.lock", "requirements-macos-arm64.lock"):
        text = (root / f).read_text()
        assert "websockets" in text and "uvicorn" in text, f


def test_recorder_flags_blocked_microphone():
    import numpy as np

    from mindbridge.audio.io import Recorder

    r = Recorder()
    r._chunks = [np.zeros(16000, np.float32)]
    r.stop()
    assert r.last_was_digital_silence
    r._chunks = [np.full(16000, 1e-4, np.float32)]
    r.stop()
    assert not r.last_was_digital_silence
